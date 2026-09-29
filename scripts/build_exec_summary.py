"""Build Word version of the Executive Summary for Stakeholders."""
from pathlib import Path
import docx

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "reports" / "Executive_Summary_Stakeholders.md"
DOCX_PATH = ROOT / "reports" / "Executive_Summary_Stakeholders.docx"

def build():
    doc = docx.Document()
    doc.add_heading("Executive Summary & Policy Briefing", level=0)
    doc.add_heading("System Capacity & Care Load Analytics for Unaccompanied Children (UAC)", level=1)
    
    content = MD_PATH.read_text(encoding="utf-8")
    for block in content.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("### "):
            doc.add_heading(block.replace("### ", ""), level=2)
        elif block.startswith("## "):
            doc.add_heading(block.replace("## ", ""), level=1)
        elif block.startswith("# "):
            pass
        elif block.startswith("---"):
            continue
        elif block.startswith("|"):
            continue
        else:
            doc.add_paragraph(block)
            
    doc.save(DOCX_PATH)
    print(f"Generated {DOCX_PATH}")

if __name__ == "__main__":
    build()
