"""Build the editable internship report from generated research evidence."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]


def build_report():
    r=json.loads((ROOT/'reports/research_results.json').read_text());m=r['models'];s=r['summary'];a=r['source_audit']
    selected=next(x for x in m['scores'] if x['model']==m['selected_model'])
    doc=Document();sec=doc.sections[0]
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.75)
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    for style,size in [('Title',32),('Heading 1',22),('Heading 2',13)]:
        doc.styles[style].font.name='Calibri';doc.styles[style].font.size=Pt(size);doc.styles[style].font.color.rgb=RGBColor.from_string('203A5B')
    # Remove template paragraph borders; hierarchy comes from type and spacing.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.core_properties.title='UAC Insight: Care Load Dynamics and Seven-Day Forecasting'
    doc.core_properties.author='Harish V J'
    footer=sec.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run('UAC Insight  |  ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('8392A4')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def p(text,style=None):return doc.add_paragraph(text,style)
    def heading(text):doc.add_heading(text,1)
    def sub(text):doc.add_heading(text,2)
    def page(label,title):doc.add_page_break();p(label,'Subtitle');heading(title)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
        for row in rows:
            cells=t.add_row().cells
            for i,value in enumerate(row):cells[i].text=str(value)
        for row in t.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(4)
                    for run in paragraph.runs:run.font.size=Pt(9)
        p('')
    def fig(name,caption,width=6.6):
        doc.add_picture(str(ROOT/f'reports/figures/{name}.png'),width=Inches(width))
        q=p(caption);q.paragraph_format.space_after=Pt(10)
        for run in q.runs:run.italic=True;run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('65778E')
    p('DATA SCIENCE INTERNSHIP PROJECT','Subtitle')
    doc.add_heading('UAC Insight',0)
    p('Care Load Dynamics and\nSeven-Day Forecasting','Title')
    p('Prepared by Harish V J\nGITAM School of Technology, Bengaluru\nSeptember 2026')
    sub('Project purpose')
    p('Transform aggregate public-program observations into an auditable analytical workflow that explains care-load patterns and tests whether forecasting adds measurable value over a simple baseline.')
    table(['Evidence','Result'],[['Source observations','720, spanning 12 Jan 2023 to 21 Dec 2025'],['Observed-date coverage','66.98%; 355 calendar dates unobserved'],['Selected forecast model',m['selected_model']],['Test mean absolute error',f"{selected['test']['mae']:.2f} children across {m['split']['test_pairs']} target pairs"],['Improvement over persistence',f"{m['test_skill_vs_persistence_pct']:.2f}% lower test MAE on this holdout"]])
    sub('Abstract')
    p('This project studies reported Customs and Border Protection (CBP) custody and Health and Human Services (HHS) care counts through a reproducible Python pipeline. It audits blank records and reporting gaps, separates stock measures from flows, engineers transparent pressure indicators, and evaluates four forecasting candidates using an exact seven-calendar-day target. A chronological split protects the test period from model selection. The selected local linear trend improves on persistence in this holdout, while more complex models do not win validation. The application communicates observed facts, project-defined indicators, model results, and limitations together.')
    sub('Why the project matters')
    p('Its contribution is a traceable path from raw records to validated evidence. Reviewers can reproduce every headline, inspect assumptions, and determine where the data stops supporting a claim. The project demonstrates data preparation, exploratory analysis, feature engineering, model evaluation, and communication in one working system.')

    page('01 / DATA UNDERSTANDING','Source, schema and quality')
    p('The source is the CSV supplied with the project, labeled HHS Unaccompanied Alien Children Program. Its latest observation is 21 December 2025; the application is a historical analysis rather than a live data feed. Source ownership and publication metadata should be confirmed before external operational deployment.')
    table(['Source audit','Count'],[['Physical rows after header',a['source_rows']],['Wholly blank rows removed',a['blank_rows_removed']],['Validated observations',a['valid_rows']],['Duplicate dates / invalid retained rows','0 / 0'],['Calendar span / unobserved dates',f"{s['calendar_days']} / {s['unobserved_days']}"],['Gaps longer than one day / largest interval',f"{s['gaps_gt_one']} / {s['largest_gap']} days"]])
    table(['Canonical field','Interpretation'],[['date','Reported observation date'],['cbp_intake','Children apprehended and placed in CBP custody'],['cbp_custody','Reported CBP custody stock'],['transfers','Children transferred out of CBP custody'],['hhs_care','Reported HHS care stock'],['discharges','Children discharged from HHS care']])
    sub('Cleaning decisions')
    p('The parser removes only rows where every field is blank. It uses the explicit month-name date format, strips thousands separators from counts, validates non-negative integers, enforces unique dates, and sorts ascending. It fails on malformed records rather than inventing replacements. Missing calendar dates remain absent; there is no silent interpolation or zero filling.')
    sub('Flow-versus-stock checks')
    p(f"There are {s['transfer_review_count']} observations with transfers above the same-date CBP custody count, and {s['discharge_review_count']} with discharges above HHS care. Different within-day measurement timing can explain such comparisons, so they are retained as review flags. They are not automatically classified as data errors.")
    sub('Provenance')
    p('The raw CSV is immutable. The run manifest records its SHA-256 checksum, the canonical-data fingerprint, package versions, and the model seed. This identifies the input that produced the report, even if the editable dashboard working copy changes later.')

    page('02 / EXPLORATORY ANALYSIS','Load changes and reporting coverage')
    fig('load_history','Figure 1. Every observed date is plotted in calendar time. Lines connect observations; the rolling mean uses seven reports.')
    p(f"The latest reported combined load is {s['latest_load']:,}, compared with a peak of {s['peak_load']:,} on {s['peak_date']}. The latest snapshot is {abs(s['change_from_peak_pct']):.2f}% below that observed peak. This is a comparison of reported headcounts and does not establish a policy effect or a capacity utilization rate.")
    table(['Year','Reports','Mean total load','Minimum','Maximum'],[[x['date'],x['observations'],f"{x['mean_load']:,.1f}",x['minimum_load'],x['maximum_load']] for x in r['yearly']])
    sub('Interpretation')
    p('The observed mean load is lower in 2025 than in the earlier years. Annual summaries are averages of available reports; uneven date coverage and partial first/last years limit direct calendar-year comparison. Monthly stock summaries also use means, whereas flows are summed only over the rows actually reported.')
    p('Coverage is a central analytical variable: 720 of 1,075 calendar dates are observed. A seven-observation rolling window can therefore span more than seven days. The dashboard labels the unit explicitly and exposes the monthly coverage calendar.')

    page('03 / FEATURE ENGINEERING','Transparent operational indicators')
    fig('flow_balance','Figure 2. Transfers, discharges and the smoothed net-flow proxy reveal different reporting regimes.',6.3)
    table(['Indicator','Definition and boundary'],[['Total load','CBP custody + HHS care; a stock, not bed utilization.'],['Net flow proxy','Transfers - discharges; not the exact census change.'],['Discharge offset','Aggregate discharges / aggregate transfers in the stated window.'],['Rolling mean','Trailing 7 or 14 reported observations, with full-window requirements.'],['Volatility','100 × SD of 14 load changes / mean load over 14 observations.'],['Pressure persistence','Positive 7-observation net sum and at least 4 positive observations.']])
    p('A flow balance supports monitoring of incoming-versus-outgoing activity. It cannot reconstruct a complete conservation equation without aligned timestamps, complete reporting and all relevant transitions. Zero transfer denominators yield a missing ratio, not infinity or a fabricated neutral value.')

    page('04 / SENSITIVITY ANALYSIS','How much do threshold choices matter?')
    fig('pressure_sensitivity','Figure 3. The share of flagged observations changes when the persistence convention changes.',6.3)
    table(['Rule','Flagged observations','Eligible windows','Share'],[[x['rule'],x['flagged'],x['eligible'],f"{x['share_pct']:.2f}%"] for x in r['sensitivity']])
    sub('Finding')
    p('The primary 4-of-7 rule flags 218 of 714 eligible observations (30.53%). Requiring five positive observations reduces that share to 24.09%. This difference demonstrates that the label depends partly on the analytical convention. A review should examine the continuous net-flow series alongside the flag.')
    sub('Relative-stress diagnostic')
    p(f"Four components are also exposed: load above its prior-history 90th percentile, volatility above its prior-history 90th percentile, positive recent mean net flow, and offset below one. Thresholds use only preceding observations. Of {s['eligible_stress_rows']} eligible records, {s['high_relative_stress_rows']} meet at least three components. This conservative composite provides no useful high-alert separation on this dataset; it is a diagnostic result, not evidence that operational risk was absent.")
    sub('What would strengthen this measure?')
    p('A useful operational alert requires an agreed target outcome and external validation. Facility capacity, staffing, reporting timestamps and placement delays would allow thresholds to be checked against real outcomes. Without those data, the indicators remain descriptive research features.')

    page('05 / MODELING PROTOCOL','Forecast design and leakage prevention')
    sub('Prediction task')
    p('Predict HHS care exactly seven calendar days after an observed origin. A training pair exists only if that target date is reported. No imputed future values enter the target. Models receive only current and earlier observations; the dataset does not contain individual-child outcomes.')
    table(['Partition','Boundary and valid pairs'],[['Training',f"Targets before {m['split']['validation_start']}; {m['split']['train_pairs']} pairs"],['Validation',f"Origins from {m['split']['validation_start']}, targets before {m['split']['test_start']}; {m['split']['validation_pairs']} pairs"],['Test',f"Origins from {m['split']['test_start']}; {m['split']['test_pairs']} pairs"]])
    p('Boundaries correspond to the 70%, 85% and end positions of the original chronological observations. Pairs crossing a boundary are excluded. After validation selects the model, learned candidates are refitted on all eligible pairs whose targets occur before the test boundary. Test scores are computed afterward and never determine selection.')
    table(['Candidate','Fixed specification'],[['Persistence','Future census equals the origin census.'],['Local linear trend','Calendar-time slope over 14 observations; extrapolate 7 days from current census.'],['Ridge regression','Standardized predictors; alpha = 10; predict census change.'],['Random forest','160 trees; depth 6; minimum leaf 8; seed 42; predict census change.']])
    sub('Feature set')
    p('Predictors include current care, custody, intake, transfers and discharges; previous care; 7/14-observation care means; 14-observation care standard deviation; recent net flow; elapsed reporting gap; care change per elapsed day; 14-observation calendar span; and the local calendar-time slope. Ridge standardization is fitted inside the training pipeline.')
    sub('Evaluation choice')
    p('MAE is the primary score because it is expressed directly in children. RMSE emphasizes larger misses; signed bias exposes systematic over- or underprediction. Percentage accuracy is not used for this regression task. Candidate settings are fixed rather than tuned on the holdout.')

    page('06 / EXPERIMENT RESULTS','A simpler model wins validation')
    fig('model_comparison','Figure 4. All candidates are evaluated on identical exact-date pairs.',6.6)
    table(['Model','Validation MAE','Test MAE','Test RMSE'],[[x['model'],f"{x['validation']['mae']:.2f}",f"{x['test']['mae']:.2f}",f"{x['test']['rmse']:.2f}"] for x in m['scores']])
    p(f"Validation selects {m['selected_model'].lower()}, with validation MAE {selected['validation']['mae']:.2f}. On the untouched test period its MAE is {selected['test']['mae']:.2f}, compared with {m['scores'][0]['test']['mae']:.2f} for persistence. This corresponds to {m['test_skill_vs_persistence_pct']:.2f}% lower MAE on this holdout.")
    sub('Interpretation and limits')
    p('The result supports retaining the simple local trend for this experiment. Random forest performs poorly on validation and does not qualify as the selected model even though its later test performance improves. Different period-level errors suggest temporal instability; this is an interpretation of the observed scores, not proof of a particular cause. More complex models need additional evidence to justify use.')
    p('This is a single historical holdout, not a universal model ranking. Nearby origins overlap and forecast errors are dependent. Further rolling-origin studies across more regimes are required before operational deployment.')

    page('07 / ERROR ANALYSIS','Inspect predictions, then bound the claim')
    fig('test_predictions','Figure 5. Predictions are aligned to the reported target date. Each origin is exactly seven days earlier.')
    p(f"The selected model has test RMSE {selected['test']['rmse']:.2f} and signed bias {selected['test']['bias']:.2f} children. The negative bias means it underpredicts on average in this test period. Per-origin predictions are included in reports/tables/test_predictions.csv so each miss can be inspected.")
    f=m['forecast'];sub('Estimate at the historical cutoff')
    p(f"Using observations available through {f['origin']}, the selected model estimates approximately {f['estimate']:,} children in HHS care on {f['date']}. This date lies beyond the supplied source cutoff and is not a present-day prediction.")
    p(f"The displayed reference range is {f['error_reference_low']:,} to {f['error_reference_high']:,}, constructed by adding and subtracting the 90th percentile of absolute validation errors ({f['error_radius']:.1f}). It is descriptive; temporal dependence and regime changes prevent treating it as a calibrated 90% prediction interval.")
    sub('Material limitations')
    p('Exact-date target availability selects a subset of observations and may introduce selection bias. No missing-date mechanism is modeled. The data are aggregate, observation times are not fully specified, and no bed or staffing capacity is provided. These results therefore support historical description and a bounded forecasting experiment, not causal explanations, individual risk estimates or verified overcrowding claims.')

    page('08 / IMPLEMENTATION & CONCLUSION','A submission that can be reproduced')
    table(['Layer','Implementation and responsibility'],[['Source and audit','src/data.py: schema validation, immutable CSV, fingerprints.'],['Features','src/features.py: interpretable metrics and origin-safe predictors.'],['Models','src/models.py: exact targets, temporal split, baselines, learned models.'],['Evidence','src/analysis.py and scripts/: tables, figures, notebook and report.'],['Application','app.py REST API, store.py SQLite copy, web/ dashboard.'],['Verification','tests/: source checks, temporal boundaries, metric edge cases and CRUD.']])
    sub('Reproduction commands')
    for command in ['py -3.12 -m venv .venv', '.venv\\Scripts\\python.exe -m pip install -r requirements.txt', '.venv\\Scripts\\python.exe scripts/build_project.py', '.venv\\Scripts\\python.exe scripts/build_notebook.py', '.venv\\Scripts\\python.exe scripts/build_report.py', '.venv\\Scripts\\python.exe -m unittest discover -s tests -v', '.venv\\Scripts\\python.exe app.py']:
        q=p(command);q.paragraph_format.space_after=Pt(3)
        for run in q.runs:run.font.name='Consolas';run.font.size=Pt(9)
    p('The dashboard opens at http://127.0.0.1:8501. All report artifacts derive from the original CSV; editing the local database changes dashboard results and triggers a visible source-mismatch notice.')
    sub('Conclusion')
    p('The project demonstrates an end-to-end Data Science workflow with a measurable result: an interpretable local trend lowers test MAE relative to persistence within the defined experiment. Equally important, the source audit and sensitivity analysis show where apparent precision would be misleading. The deliverable makes these boundaries inspectable through shared code, generated evidence and an interactive application.')
    sub('Sources and reproducible evidence')
    p('[1] Supplied HHS_Unaccompanied_Alien_Children_Program.csv, preserved in data/. Coverage: 12 January 2023 to 21 December 2025. Exact checksum is recorded in reports/run_manifest.json.')
    p('[2] This project: src/data.py, src/features.py, src/models.py and src/analysis.py. Formula definitions and fixed candidate configurations used for this experiment.')
    p('[3] Generated evidence: reports/research_results.json, reports/tables/, reports/figures/ and notebooks/01_research_workflow.ipynb.')
    output=ROOT/'reports/Internship_Report.docx';doc.save(output);print(output)


if __name__=='__main__':build_report()
